"""DOCX generation for CHART preparedness plans.

Follows the Muzaffarpur DHAP template structure with 6 sections.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .models import (
    PHASE_MW_RANGES,
    SEVERITY_LABELS,
    MatchedIntervention,
    Phase,
    PreparednesssPlan,
    Severity,
)

# Color mapping for severity levels
SEVERITY_COLORS: dict[Severity, RGBColor] = {
    Severity.GREEN: RGBColor(0x00, 0xB0, 0x50),
    Severity.YELLOW: RGBColor(0xFF, 0xD6, 0x00),
    Severity.ORANGE: RGBColor(0xFF, 0x8C, 0x00),
    Severity.RED: RGBColor(0xFF, 0x00, 0x00),
}

PHASE_LABELS: dict[Phase, str] = {
    Phase.PRE_MONSOON: "Pre-Monsoon",
    Phase.ACTIVE_MONSOON: "Active Monsoon",
    Phase.POST_MONSOON: "Post-Monsoon",
}

# Month names mapped to approximate MW ranges for the heatmap
MW_TO_MONTH: list[tuple[str, int, int]] = [
    ("May", 18, 22),
    ("Jun", 23, 26),
    ("Jul", 27, 31),
    ("Aug", 31, 35),
    ("Sep", 36, 39),
    ("Oct", 40, 44),
]


def _set_cell_shading(cell, color_hex: str) -> None:
    """Set background shading on a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {
            qn("w:fill"): color_hex,
            qn("w:val"): "clear",
        },
    )
    shading.append(shading_elem)


def _add_styled_paragraph(
    doc: Document,
    text: str,
    style: str = "Normal",
    bold: bool = False,
    size: int | None = None,
    alignment: int | None = None,
    space_after: int | None = None,
) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)


def _add_cover_page(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 0: Cover page."""
    profile = plan.district_profile

    for _ in range(3):
        doc.add_paragraph()

    _add_styled_paragraph(
        doc,
        f"Government of {profile.state}",
        bold=True,
        size=16,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_styled_paragraph(
        doc,
        f"District Health Society, {profile.name}",
        bold=True,
        size=14,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_paragraph()

    _add_styled_paragraph(
        doc,
        "District Health Action Plan",
        bold=True,
        size=22,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_styled_paragraph(
        doc,
        "Climate-Health Preparedness Component",
        bold=True,
        size=14,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_paragraph()

    _add_styled_paragraph(
        doc,
        f"Valid Period: {plan.valid_period}",
        size=12,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_styled_paragraph(
        doc,
        f"Plan Reference: {plan.auto_plan_reference}",
        size=12,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_styled_paragraph(
        doc,
        f"Generated: {plan.generated_at.strftime('%d %B %Y')}",
        size=12,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    for _ in range(3):
        doc.add_paragraph()

    _add_styled_paragraph(
        doc,
        "PUBLIC DOCUMENT",
        bold=True,
        size=14,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_styled_paragraph(
        doc,
        "This plan belongs to the district and its communities.",
        size=10,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    doc.add_page_break()


def _add_section1_situation_assessment(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 1: Climate-health situation assessment."""
    doc.add_heading("1. Climate-Health Situation Assessment", level=1)

    profile = plan.district_profile
    doc.add_paragraph(
        f"District {profile.name}, {profile.state} has a population of "
        f"{profile.population:,} across {len(profile.blocks)} blocks with "
        f"{profile.health_facilities.phcs} PHCs and "
        f"{profile.health_facilities.scs} Sub-Centres."
    )

    if profile.priority_populations:
        doc.add_heading("Priority Populations", level=2)
        for pp in profile.priority_populations:
            doc.add_paragraph(pp, style="List Bullet")

    # Hazard profile table
    doc.add_heading("Active Hazard Profile", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Hazard", "Severity", "Peak Period (MW)", "Health Impacts", "Confidence"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True if cell.paragraphs[0].runs else None
        # Bold the header text
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True

    for hazard in plan.hazards:
        row = table.add_row()
        row.cells[0].text = hazard.hazard_type.value.replace("_", " ").title()

        severity_label = SEVERITY_LABELS[hazard.severity]
        row.cells[1].text = severity_label
        color = SEVERITY_COLORS[hazard.severity]
        _set_cell_shading(row.cells[1], str(color))

        row.cells[2].text = f"MW {hazard.peak_period_mw_start}–{hazard.peak_period_mw_end}"
        row.cells[3].text = "; ".join(hazard.health_impacts)
        row.cells[4].text = hazard.confidence.value.title()

    doc.add_page_break()


def _add_section2_intervention_plan(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 2: Intervention plan by phase."""
    doc.add_heading("2. Intervention Plan", level=1)

    by_phase = plan.interventions_by_phase()

    for phase in Phase:
        interventions = by_phase[phase]
        mw_start, mw_end = PHASE_MW_RANGES[phase]

        doc.add_heading(
            f"{PHASE_LABELS[phase]} (MW {mw_start}–{mw_end})",
            level=2,
        )

        if not interventions:
            doc.add_paragraph("No interventions matched for this phase.")
            continue

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["#", "Intervention", "Trigger", "Responsible", "Indicator"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True

        for mi in interventions:
            row = table.add_row()
            row.cells[0].text = str(mi.sequence_number)
            row.cells[1].text = mi.intervention.name
            row.cells[2].text = mi.intervention.trigger_condition
            row.cells[3].text = mi.intervention.responsible_officer_role
            row.cells[4].text = mi.intervention.indicator

    doc.add_page_break()


def _add_section3_timeline(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 3: Implementation timeline — hazard heatmap by month."""
    doc.add_heading("3. Implementation Timeline", level=1)
    doc.add_paragraph(
        "The following heatmap shows hazard activity across the planning period. "
        "Shaded cells indicate months where the hazard is expected to be active."
    )

    months = [m[0] for m in MW_TO_MONTH]
    table = doc.add_table(rows=1, cols=len(months) + 1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    table.rows[0].cells[0].text = "Hazard"
    for p in table.rows[0].cells[0].paragraphs:
        for run in p.runs:
            run.bold = True
    for i, month in enumerate(months):
        table.rows[0].cells[i + 1].text = month
        for p in table.rows[0].cells[i + 1].paragraphs:
            for run in p.runs:
                run.bold = True

    for hazard in plan.hazards:
        row = table.add_row()
        row.cells[0].text = hazard.hazard_type.value.replace("_", " ").title()

        for i, (_, mw_start, mw_end) in enumerate(MW_TO_MONTH):
            # Check overlap between hazard peak and this month's MW range
            if hazard.peak_period_mw_start <= mw_end and hazard.peak_period_mw_end >= mw_start:
                color = SEVERITY_COLORS[hazard.severity]
                _set_cell_shading(row.cells[i + 1], str(color))
                row.cells[i + 1].text = SEVERITY_LABELS[hazard.severity]

    doc.add_page_break()


def _add_section4_monitoring(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 4: Monitoring framework."""
    doc.add_heading("4. Monitoring Framework", level=1)

    if not plan.monitoring_indicators:
        doc.add_paragraph("No monitoring indicators defined.")
        return

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Indicator", "Baseline", "Target", "Reporting Frequency"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True

    for ind in plan.monitoring_indicators:
        row = table.add_row()
        row.cells[0].text = ind.indicator
        row.cells[1].text = ind.baseline
        row.cells[2].text = ind.target
        row.cells[3].text = ind.reporting_frequency

    doc.add_page_break()


def _add_section5_data_sources(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 5: Data sources and methodology."""
    doc.add_heading("5. Data Sources and Methodology", level=1)
    doc.add_paragraph(
        "This plan was generated using the CHART (Climate x Health Adaptation "
        "and Resilience Tool) plan builder. The following data sources were used:"
    )
    for source in plan.data_sources:
        doc.add_paragraph(source, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph(
        "Methodology: Hazard-intervention matching was performed against a curated "
        "knowledge base of climate-health interventions. Interventions were sequenced "
        "by seasonal phase (pre-monsoon, active monsoon, post-monsoon) and prioritized "
        "by hazard severity level."
    )


def _add_section6_revision(doc: Document, plan: PreparednesssPlan) -> None:
    """Section 6: Revision checkpoints and CMO approval."""
    doc.add_heading("6. Revision Checkpoints", level=1)

    for cp in plan.revision_checkpoints:
        doc.add_paragraph(f"{cp.date_description}: {cp.trigger}")

    # CMO approval signature block
    doc.add_paragraph()
    doc.add_heading("Approval", level=2)

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Prepared by"
    table.rows[0].cells[1].text = "District Programme Manager (NHM)"
    table.rows[1].cells[0].text = "Reviewed by"
    table.rows[1].cells[1].text = "District Health Officer"
    table.rows[2].cells[0].text = "Approved by"
    table.rows[2].cells[1].text = "Chief Medical Officer (CMO)"

    doc.add_paragraph()
    doc.add_paragraph("Signature: ____________________    Date: ____________")
    doc.add_paragraph()
    doc.add_paragraph("(CMO Seal)")


def generate_docx(plan: PreparednesssPlan, output_path: Path) -> Path:
    """Generate a DOCX document from a preparedness plan.

    Args:
        plan: The fully assembled preparedness plan.
        output_path: Where to write the .docx file.

    Returns:
        The output path (for convenience).
    """
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    _add_cover_page(doc, plan)
    _add_section1_situation_assessment(doc, plan)
    _add_section2_intervention_plan(doc, plan)
    _add_section3_timeline(doc, plan)
    _add_section4_monitoring(doc, plan)
    _add_section5_data_sources(doc, plan)
    _add_section6_revision(doc, plan)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
